"""PDF and XLSX exports for normalized dashboard reports."""
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import DoughnutChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


INK = '102A25'
GREEN = '087A58'
GREEN_SOFT = 'DFF3E9'
RED = 'B63B32'
RED_SOFT = 'F9E2DF'
AMBER = 'A05A00'
AMBER_SOFT = 'FFF0D2'
LINE = 'CBD7CF'


def _scenario_rows(report):
    for feature in report['features']:
        for scenario in feature['scenarios']:
            yield feature, scenario


def _status_fill(status):
    return {
        'passed': PatternFill('solid', fgColor=GREEN_SOFT),
        'failed': PatternFill('solid', fgColor=RED_SOFT),
        'skipped': PatternFill('solid', fgColor=AMBER_SOFT),
    }.get(status, PatternFill('solid', fgColor='EFF3F0'))


def _style_header(sheet, row, end_column):
    for column in range(1, end_column + 1):
        cell = sheet.cell(row, column)
        cell.fill = PatternFill('solid', fgColor=INK)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[row].height = 24


def _size_columns(sheet):
    for column_index, column in enumerate(sheet.columns, 1):
        width = min(max(len(str(cell.value or '')) for cell in column) + 2, 52)
        sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)


def _write_table(sheet, start_row, headers, rows, status_column=None):
    for column, header in enumerate(headers, 1):
        sheet.cell(start_row, column, header)
    _style_header(sheet, start_row, len(headers))
    for row_index, values in enumerate(rows, start_row + 1):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=Side(style='hair', color=LINE))
        if status_column:
            status_cell = sheet.cell(row_index, status_column)
            status_cell.fill = _status_fill(status_cell.value)
            status_cell.font = Font(bold=True, color=INK)
    if rows:
        sheet.auto_filter.ref = f'A{start_row}:{sheet.cell(start_row + len(rows), len(headers)).coordinate}'
    sheet.freeze_panes = f'A{start_row + 1}'


def build_xlsx(report):
    """Build an executive XLSX workbook from a normalized dashboard report."""
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = 'Resumo'
    summary_sheet.merge_cells('A1:F1')
    summary_sheet['A1'] = 'AutomacaoBDD | Relatorio de resultados'
    summary_sheet['A1'].fill = PatternFill('solid', fgColor=INK)
    summary_sheet['A1'].font = Font(bold=True, color='FFFFFF', size=16)
    summary_sheet['A1'].alignment = Alignment(vertical='center')
    summary_sheet.row_dimensions[1].height = 34
    summary_sheet.merge_cells('A2:F2')
    summary_sheet['A2'] = (
        f"{report['meta']['application']} | ambiente: {report['meta']['environment'] or 'nao informado'} | "
        f"versao: {report['meta']['version'] or 'nao informada'} | {report['meta']['date']}"
    )
    summary_sheet['A2'].font = Font(color='60716A', italic=True)
    summary_sheet['A4'] = 'STATUS DA EXECUCAO'
    summary_sheet['B4'] = 'APROVADA' if not report['failures'] and report['summary']['scenarios'] else 'ATENCAO'
    summary_sheet['B4'].fill = _status_fill('passed' if summary_sheet['B4'].value == 'APROVADA' else 'failed')
    summary_sheet['B4'].font = Font(bold=True)
    summary_sheet['A6'] = 'INDICADOR'
    summary_sheet['B6'] = 'VALOR'
    _style_header(summary_sheet, 6, 2)
    for row_index, (label, value) in enumerate([
        ('Features executadas', report['summary']['features']),
        ('Cenarios avaliados', report['summary']['scenarios']),
        ('Cenarios aprovados', report['summary']['passed']),
        ('Cenarios com erro', report['summary']['failed']),
        ('Steps executados', report['summary']['steps']),
        ('Taxa de aprovacao', f"{report['analytics']['pass_rate']}%"),
        ('Duracao total', f"{report['meta']['duration']}s"),
        ('Media por cenario', f"{report['analytics']['average_scenario_duration']}s"),
    ], 7):
        summary_sheet.cell(row_index, 1, label)
        summary_sheet.cell(row_index, 2, value)
        summary_sheet.cell(row_index, 1).border = Border(bottom=Side(style='hair', color=LINE))
        summary_sheet.cell(row_index, 2).border = Border(bottom=Side(style='hair', color=LINE))
    summary_sheet['D6'] = 'STATUS DE STEPS'
    summary_sheet['E6'] = 'TOTAL'
    _style_header(summary_sheet, 6, 5)
    for row_index, status in enumerate(('passed', 'failed', 'skipped'), 7):
        summary_sheet.cell(row_index, 4, status.upper())
        summary_sheet.cell(row_index, 5, report['analytics']['step_statuses'][status])
        summary_sheet.cell(row_index, 4).fill = _status_fill(status)
    chart = DoughnutChart()
    chart.title = 'Cenarios por status'
    chart.add_data(Reference(summary_sheet, min_col=2, min_row=9, max_row=10))
    chart.set_categories(Reference(summary_sheet, min_col=1, min_row=9, max_row=10))
    chart.height = 6
    chart.width = 10
    summary_sheet.add_chart(chart, 'D10')
    summary_sheet.freeze_panes = 'A6'

    scenarios_sheet = workbook.create_sheet('Cenarios')
    scenario_rows = []
    for feature, scenario in _scenario_rows(report):
        scenario_rows.append([
            feature['name'], scenario['name'], scenario['status'], scenario['duration'], scenario['location'],
        ])
    _write_table(scenarios_sheet, 1, ['Feature', 'Cenario', 'Status', 'Duracao (s)', 'Localizacao'], scenario_rows, 3)

    steps_sheet = workbook.create_sheet('Steps')
    step_rows = []
    for feature, scenario in _scenario_rows(report):
        for step in scenario['steps']:
            step_rows.append([
                feature['name'], scenario['name'], f"{step['keyword']} {step['name']}", step['status'],
                step['duration'], step['expected'], step['actual'],
            ])
    _write_table(steps_sheet, 1, ['Feature', 'Cenario', 'Step', 'Status', 'Duracao (s)', 'Esperado', 'Obtido'], step_rows, 4)

    failures_sheet = workbook.create_sheet('Falhas')
    failure_rows = []
    for failure in report['failures']:
        step = failure['step'] or {}
        failure_rows.append([
            failure['feature'], failure['scenario'], step.get('name', 'Cenario falhou sem step identificado'),
            step.get('actual', ''),
        ])
    _write_table(failures_sheet, 1, ['Feature', 'Cenario', 'Ponto de falha', 'Detalhe obtido'], failure_rows)
    if not failure_rows:
        failures_sheet['A2'] = 'Nenhuma falha registrada nesta execucao.'
        failures_sheet['A2'].font = Font(italic=True, color='60716A')

    for sheet in workbook.worksheets:
        _size_columns(sheet)
        sheet.sheet_view.showGridLines = False

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _pdf_footer(canvas, document):
    canvas.saveState()
    canvas.setFont('Helvetica', 8)
    canvas.setFillColor(colors.HexColor('#60716A'))
    canvas.drawString(document.leftMargin, 1.1 * cm, 'AutomacaoBDD | Relatorio gerado automaticamente')
    canvas.drawRightString(A4[0] - document.rightMargin, 1.1 * cm, f'Pagina {document.page}')
    canvas.restoreState()


def build_pdf(report):
    """Build a readable PDF executive report from dashboard data."""
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.35 * cm, bottomMargin=1.75 * cm,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('ReportTitle', parent=styles['Title'], textColor=colors.HexColor('#102A25'), spaceAfter=4))
    styles.add(ParagraphStyle('ReportMeta', parent=styles['Normal'], textColor=colors.HexColor('#60716A'), fontSize=9))
    styles.add(ParagraphStyle('ReportBody', parent=styles['Normal'], textColor=colors.HexColor('#31443D'), fontSize=9, leading=13))
    healthy = report['summary']['scenarios'] and not report['failures']
    status_label = 'SUITE APROVADA' if healthy else 'EXECUCAO COM ATENCAO'
    status_color = '#087A58' if healthy else '#B63B32'
    story = [Paragraph('AutomacaoBDD | Relatorio de resultados', styles['ReportTitle'])]
    story.append(Paragraph(
        escape(f"{report['meta']['application']} | {report['meta']['environment']} | versao {report['meta']['version'] or 'nao informada'} | {report['meta']['date']}"),
        styles['ReportMeta'],
    ))
    story.append(Spacer(1, 12))
    status_table = Table([[status_label, f"{report['analytics']['pass_rate']}% de aprovacao", f"{report['meta']['duration']}s de duracao"]], colWidths=[6 * cm, 6 * cm, 6 * cm])
    status_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, 0), colors.HexColor(status_color)),
        ('BACKGROUND', (1, 0), (-1, -1), colors.HexColor('#E9F0EB')),
        ('TEXTCOLOR', (0, 0), (0, 0), colors.white), ('TEXTCOLOR', (1, 0), (-1, -1), colors.HexColor('#102A25')),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'), ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 10), ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.extend([status_table, Spacer(1, 14), Paragraph('Resumo executivo', styles['Heading2'])])
    summary_data = [
        ['Features', 'Cenarios', 'Passaram', 'Erros', 'Taxa de aprovacao', 'Duracao'],
        [
            report['summary']['features'], report['summary']['scenarios'], report['summary']['passed'],
            report['summary']['failed'], f"{report['analytics']['pass_rate']}%", f"{report['meta']['duration']}s",
        ],
    ]
    summary_table = Table(summary_data, colWidths=[2.6 * cm] * 6)
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#102A25')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'), ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#C9D4CE')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))
    slowest = report['analytics']['slowest_scenario']
    analysis = f"Media de {report['analytics']['average_scenario_duration']}s por cenario. "
    analysis += f"Cenario mais lento: {slowest['name']} ({slowest['duration']}s)." if slowest else 'Nenhum cenario disponivel para analise.'
    story.extend([summary_table, Spacer(1, 12), Paragraph(escape(analysis), styles['ReportBody']), Spacer(1, 14), Paragraph('Cenarios', styles['Heading2'])])
    scenario_data = [['Feature', 'Cenario', 'Status', 'Duracao']]
    for feature, scenario in _scenario_rows(report):
        scenario_data.append([
            Paragraph(escape(feature['name']), styles['ReportBody']), Paragraph(escape(scenario['name']), styles['ReportBody']),
            scenario['status'].upper(), f"{scenario['duration']}s",
        ])
    if len(scenario_data) == 1:
        scenario_data.append(['-', 'Nenhum cenario registrado', '-', '-'])
    scenario_table = Table(scenario_data, colWidths=[5 * cm, 7 * cm, 2.1 * cm, 2.1 * cm], repeatRows=1)
    scenario_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E5EEE8')),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#C9D4CE')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(scenario_table)
    if report['failures']:
        story.extend([Spacer(1, 14), Paragraph('Falhas que requerem analise', styles['Heading2'])])
        failure_data = [['Cenario', 'Ponto de falha']]
        for failure in report['failures']:
            step = failure['step'] or {}
            detail = step.get('actual') or step.get('name') or 'Cenario falhou sem detalhe adicional.'
            failure_data.append([Paragraph(escape(failure['scenario']), styles['ReportBody']), Paragraph(escape(detail), styles['ReportBody'])])
        failure_table = Table(failure_data, colWidths=[7 * cm, 9.2 * cm], repeatRows=1)
        failure_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F9E2DF')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#772B26')),
            ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#E8BBB6')), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ]))
        story.append(failure_table)
    document.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    output.seek(0)
    return output