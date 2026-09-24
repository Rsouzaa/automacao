from io import BytesIO

from openpyxl import load_workbook

from dashboard.app import create_app
from dashboard.report_service import load_report, normalize_report


def test_normalize_report_builds_summary_and_failure_details():
    report = {
        'features': [{'name': 'Login', 'status': 'failed', 'duration': 3.4, 'elements': [{
            'name': 'Invalid password', 'status': 'failed', 'duration': 1.2, 'location': 'login.feature:4',
            'steps': [{'keyword': 'Then', 'name': 'show error', 'result': {
                'status': 'failed', 'duration': 0.2, 'obtained_result': 'Message not displayed',
            }}],
        }]}],
        'global_data': {'application': 'Portal', 'results': {
            'total_features': 1, 'total_scenarios': 1, 'total_steps': 1,
            'passed_scenarios': 0, 'failed_scenarios': 1, 'steps_skipped': 0,
            'start_time': 10, 'end_time': 13,
        }},
    }

    dashboard = normalize_report(report)

    assert dashboard['summary']['failed'] == 1
    assert dashboard['meta']['duration'] == 3
    assert dashboard['failures'][0]['step']['actual'] == 'Message not displayed'
    assert dashboard['analytics']['pass_rate'] == 0


def test_load_report_returns_empty_state_when_report_is_missing(tmp_path):
    dashboard = load_report(tmp_path / 'missing.json')

    assert dashboard['summary']['features'] == 0
    assert 'Report not found' in dashboard['message']


def test_dashboard_exports_pdf_and_xlsx(tmp_path):
    report_path = tmp_path / 'report.json'
    report_path.write_text('{"features": [], "global_data": {"results": {}}}', encoding='utf-8')
    client = create_app(report_path).test_client()

    pdf_response = client.get('/api/export/pdf')
    xlsx_response = client.get('/api/export/xlsx')

    assert pdf_response.status_code == 200
    assert pdf_response.data.startswith(b'%PDF')
    assert xlsx_response.status_code == 200
    assert xlsx_response.data.startswith(b'PK')
    workbook = load_workbook(BytesIO(xlsx_response.data))
    assert workbook.sheetnames == ['Resumo', 'Cenarios', 'Steps', 'Falhas']
    assert workbook['Resumo']['A1'].value == 'AutomacaoBDD | Relatorio de resultados'